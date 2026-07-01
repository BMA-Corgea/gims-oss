# api/iostore/loaders.py -- split out of api/i_o.py (wiring-neutral). Generic load + file open.
from __future__ import annotations
import io
import builtins
import json
import csv
import zipfile
import mimetypes
import re
from pathlib import Path
from api.json_proxy import read_text, write_text, S3_ENABLED, _is_s3_path
from .fs_shims import fs_stat_size
from .fs_io import fs_open_readbin
from utils.logger import get_logger

log = get_logger(__name__)


def load_data(path: Path, *, default=None, strict: bool = False, encoding: str = "utf-8"):
    """
    Read JSON (or text if .md/.csv) from `path`. On missing file in S3/FS:
      - if strict=True → raise
      - else → return `default`
    """
    try:
        text = read_text(path, encoding=encoding)  # S3-aware
    except Exception as e:
        msg = str(e)
        not_found = (
            isinstance(e, FileNotFoundError)
            or "NoSuchKey" in msg
            or "The specified key does not exist" in msg
            or "Not Found" in msg
        )
        if not strict and not_found:
            return default
        raise ValueError(f"X Failed to load data from {path}: {e}")

    low = path.suffix.lower()
    if low == ".json":
        return json.loads(text) if text else default
    if low == ".jsonl":
        # Return list for JSONL
        return [json.loads(line) for line in text.splitlines() if line.strip()] if text else default
    if low in {".md", ".csv", ".txt", ""}:
        return text if text is not None else default
    return json.loads(text) if text else default

def is_file_empty(path: Path) -> bool:
    """
    S3-aware emptiness check.
    - For text-like types, read via read_text (S3-aware).
    - For CSV, stream via open_file (S3-aware).
    - For binaries (xlsx/docx/pdf/zip/images), stream via fs_open_readbin (S3-aware).
    - Falls back to size/stat checks using fs_stat_size for S3 and local stat() for FS.
    """
    ext = path.suffix.lower()

    try:
        # --- Text-ish files: read as text (works for S3 + local)
        if ext in {".txt", ".log", ".md", ".html", ".xml", ".json"}:
            return not (read_text(path, encoding="utf-8", errors="ignore").strip())

        # --- CSV: iterate rows (S3 + local via open_file)
        if ext == ".csv":
            with open_file(path, "r", encoding="utf-8") as f:
                reader = csv.reader(f)
                return not any(row for row in reader)

        # --- Quick size check (S3 uses fs_stat_size; local uses stat)
        if S3_ENABLED and _is_s3_path(path):
            if fs_stat_size(path) == 0:
                return True
        else:
            if (not path.exists()) or path.stat().st_size == 0:
                return True

        # --- XLSX: check for any non-empty cell (S3 + local)
        if ext == ".xlsx":
            import openpyxl
            with fs_open_readbin(path) as fh:
                wb = openpyxl.load_workbook(fh, data_only=True)
            return all(
                not any(cell.value for row in sheet.iter_rows() for cell in row)
                for sheet in wb.worksheets
            )

        # --- DOCX: check for any non-empty paragraph (S3 + local)
        if ext == ".docx":
            import docx
            with fs_open_readbin(path) as fh:
                doc = docx.Document(fh)
            return not any((p.text or "").strip() for p in doc.paragraphs)

        # --- PDF: any page with extractable text? (S3 + local)
        if ext == ".pdf":
            import PyPDF2
            with fs_open_readbin(path) as fh:
                reader = PyPDF2.PdfReader(fh)
                return all(not ((page.extract_text() or "").strip()) for page in reader.pages)

        # --- ZIP: all members empty or only directories? (S3 + local)
        # (We skip zipfile.is_zipfile(path) to be S3-safe and inspect via ZipFile directly.)
        try:
            with fs_open_readbin(path) as fh:
                with zipfile.ZipFile(fh) as z:
                    return all(
                        info.file_size == 0
                        for info in z.infolist()
                        if not info.filename.endswith("/")
                    )
        except zipfile.BadZipFile:
            # Not a zip; continue to image/mime handling below
            pass

        # --- Images: check for 0x0 size (S3 + local)
        if ext in {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tiff"}:
            from PIL import Image
            with fs_open_readbin(path) as fh:
                with Image.open(fh) as img:
                    return img.size == (0, 0)

        # --- Fallback mime-based text sniff
        mime_type, _ = mimetypes.guess_type(path.name)
        if mime_type and mime_type.startswith("text"):
            return not (read_text(path, encoding="utf-8", errors="ignore").strip())

        # If we got here, treat as non-empty by default (binary of unknown type)
        return False

    except FileNotFoundError:
        return True
    except Exception as e:
        log.debug(f"! Error checking {path.name}: {e}")
        return False

def _sanitize_table_name(noun: str) -> str:
    base = re.sub(r"[^0-9a-zA-Z_]", "_", noun).strip("_")
    if not base or not base[0].isalpha():
        base = f"T_{base}"
    return f"noun_{base}"

def get_url_base(project_path: Path) -> str:
    """
    Return the base API URL for this project. (S3-AWARE)
    Looks for <project>/url_base.txt. If missing, fall back to localhost:8000.
    """
    f = project_path / "url_base.txt"
    try:
        return read_text(f, encoding="utf-8").strip()
    except FileNotFoundError:
        return "http://127.0.0.1:8000"

def open_file(path, mode="r", encoding="utf-8", **kwargs):
    """
    Unified file opener for local and S3 paths.
    Accepts extra **kwargs (e.g., errors="ignore") to mirror builtins.open API.
    """
    errors = kwargs.get("errors", None)

    # ---- Local FS fast-path ---------------------------------------------------
    if not S3_ENABLED:
        if "b" in mode:
            return builtins.open(path, mode)
        return builtins.open(path, mode, encoding=encoding, errors=errors)

    # ---- S3 mode --------------------------------------------------------------
    try:
        from api import json_proxy as _jp  # type: ignore
    except Exception:
        _jp = None

    has_rb = bool(_jp and getattr(_jp, "read_bytes", None))
    has_wb = bool(_jp and getattr(_jp, "write_bytes", None))

    # BINARY MODES --------------------------------------------------------------
    if "b" in mode:
        # Read
        if "r" in mode and has_rb:
            data = _jp.read_bytes(path)  # type: ignore[attr-defined]
            return io.BytesIO(data)

        # Write / Append
        if ("w" in mode or "a" in mode) and has_wb:
            initial = b""
            if "a" in mode:
                try:
                    initial = _jp.read_bytes(path)  # type: ignore[attr-defined]
                except FileNotFoundError:
                    initial = b""

            buffer = io.BytesIO(initial)

            # Safer close monkey-patch that preserves the original close()
            orig_close = buffer.close
            def _close_and_upload_bin():
                body = buffer.getvalue()
                _jp.write_bytes(path, body)  # type: ignore[attr-defined]
                orig_close()
            buffer.close = _close_and_upload_bin  # type: ignore[assignment]

            return buffer

        # No byte helpers available → fall back to local FS (best-effort)
        log.debug("[i_o.open_file] WARNING: Binary mode without S3 byte helpers; using local FS.")
        return builtins.open(path, mode)

    # TEXT MODES ---------------------------------------------------------------
    if "r" in mode:
        # read_text is S3-aware
        data = read_text(path, encoding=encoding, errors=errors)
        return io.StringIO(data)

    if "w" in mode or "a" in mode:
        # Seed with existing content for append
        initial_data = ""
        if "a" in mode:
            try:
                initial_data = read_text(path, encoding=encoding)
                if not initial_data.endswith("\n"):
                    initial_data += "\n"
            except FileNotFoundError:
                pass

        buffer = io.StringIO(initial_data)
        if "a" in mode:
            buffer.seek(0, io.SEEK_END)

        # Safer close monkey-patch that preserves the original close()
        orig_close = buffer.close
        def _close_and_upload():
            body = buffer.getvalue()
            write_text(path, body, encoding=encoding)
            orig_close()
        buffer.close = _close_and_upload  # type: ignore[assignment]

        return buffer

    raise ValueError(f"Unsupported mode: {mode}")
